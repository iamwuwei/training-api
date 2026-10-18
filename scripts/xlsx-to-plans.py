#!/usr/bin/env python3
"""Convert a 課表 workbook into seed/plans.json (one entry per user, merged in place).

Usage:  python3 scripts/xlsx-to-plans.py [path/to/課表.xlsx] [--user ww|gf] [--goal H:MM] [--days N]
        python3 scripts/xlsx-to-plans.py                                   # ww from ../running
        python3 scripts/xlsx-to-plans.py ~/Downloads/Sara課表.xlsx --user gf --goal 5:30
Writes  seed/plans.json — replaces only the given user's entry, keeps the others.

Reads
  課表   A–I  -> sessions (plan)           J–W -> results (actuals, hand-typed verdicts)
  週總覽 A–D, I -> weeks (phase + focus)
Unknown 類型 values abort with a list so they can be added to TYPE_MAP.
"""
from __future__ import annotations

import datetime as dt
import json
import pathlib
import re
import sys

import openpyxl

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent
DEFAULT_XLSX = ROOT.parent / "running" / "富士山馬拉松_2026_課表.xlsx"  # ../running checkout
OUT = ROOT / "seed" / "plans.json"

import argparse

RACE_NAME = "富士山マラソン"
RACE_DATE = "2026-12-13"  # Sunday; the xlsx still says 12/14 on the race row
# Per-user defaults; --goal / --days override. Keys must match USERS in seed/seed.ts.
DEFAULTS = {
    "ww": {"goal_time_sec": 4 * 3600 + 45 * 60, "days_per_week": 4},  # 比賽目標 4:45
    "gf": {"goal_time_sec": None, "days_per_week": None},
}

# 類型 (column E) -> easy | long | tempo | interval | rest | race.
# The original label is kept in `label` for the UI avatars (E / R / L / MP / 坡 / 肌).
TYPE_MAP = {
    "Easy": "easy",
    "Recovery": "easy",
    "快步": "easy",          # easy run + strides
    "Long": "long",
    "Peak Long": "long",
    "MP": "tempo",           # marathon-pace segments
    "坡跑": "interval",      # hill repeats
    "坡跑+MP": "interval",
    "走跑Easy": "easy",      # run/walk intervals (Sara)
    "走跑Long": "long",
    "走跑Recovery": "easy",
    "走跑": "easy",
    "休息": "rest",
    "完全休息": "rest",
    "肌力": "rest",          # strength / Hyrox day, no running
    "比賽": "race",
}

PACE_RANGE = re.compile(r"(\d):(\d\d)\s*[〜~\-–]\s*(\d):(\d\d)")
PACE_PLUS = re.compile(r"(\d):(\d\d)\s*\+")
HR_RANGE = re.compile(r"心拍\s*(\d{2,3})\s*[〜~\-–]\s*(\d{2,3})")
HR_MAX = re.compile(r"心拍\s*[≤<=＜]+\s*(\d{2,3})")
TIME_RE = re.compile(r"^(?:(\d+):)?(\d{1,2}):(\d{2})$")


def sec(m: int, s: int) -> int:
    return m * 60 + s


def parse_target(target: str) -> dict:
    """G column -> pace min/max (sec/km), hr min/max, and a midpoint pace."""
    out = {"pace_min_sec_km": None, "pace_max_sec_km": None, "hr_min": None, "hr_max": None}
    if not target:
        out["pace_sec_km"] = None
        return out
    rng = PACE_RANGE.search(target)
    if target.startswith("上坡") and "MP" not in target:
        rng = None  # "上坡 6:00 以內" is an effort cap, not an average pace
    if rng:
        out["pace_min_sec_km"] = sec(int(rng.group(1)), int(rng.group(2)))
        out["pace_max_sec_km"] = sec(int(rng.group(3)), int(rng.group(4)))
    else:
        plus = PACE_PLUS.search(target)
        if plus:  # "7:45+" = at least this slow
            out["pace_min_sec_km"] = sec(int(plus.group(1)), int(plus.group(2)))
    hr = HR_RANGE.search(target)
    if hr:
        out["hr_min"], out["hr_max"] = int(hr.group(1)), int(hr.group(2))
    else:
        hrm = HR_MAX.search(target)
        if hrm:
            out["hr_max"] = int(hrm.group(1))
    lo, hi = out["pace_min_sec_km"], out["pace_max_sec_km"]
    out["pace_sec_km"] = round((lo + hi) / 2) if lo and hi else lo
    return out


def text(v) -> str:
    return str(v).strip() if v is not None else ""


def num(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def time_sec(v) -> int | None:
    m = TIME_RE.match(text(v))
    if not m:
        return None
    h = int(m.group(1) or 0)
    return h * 3600 + int(m.group(2)) * 60 + int(m.group(3))


def pace_sec(v) -> int | None:
    m = re.match(r"^(\d+):(\d{2})$", text(v))
    return sec(int(m.group(1)), int(m.group(2))) if m else None


def goal_sec(v: str | None) -> int | None:
    if not v:
        return None
    m = re.match(r"^(\d+):(\d{2})$", v)
    if not m:
        sys.exit(f"--goal must look like H:MM, got {v!r}")
    return int(m.group(1)) * 3600 + int(m.group(2)) * 60


def norm(h) -> str:
    """'計畫 km' -> '計畫km', '備註（補給・肌力・提醒）' -> '備註', '天氣 / 氣溫 / 濕度' -> '天氣/氣溫/濕度'."""
    return re.split(r"[（(]", re.sub(r"\s+", "", text(h)))[0]


def header_index(ws) -> dict[str, int]:
    """Normalised column name -> 0-based index from the first row, so sheets may differ in layout."""
    return {norm(h): i for i, h in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))) if text(h)}


def find_col(col: dict[str, int], name: str) -> int | None:
    """Exact normalised match first, then a header that starts with the name ('天氣' matches '天氣/氣溫/濕度')."""
    if name in col:
        return col[name]
    for h, i in col.items():
        if h.startswith(name):
            return i
    return None


def main(xlsx: pathlib.Path, user: str, goal: int | None, days: int | None) -> None:
    wb = openpyxl.load_workbook(xlsx, data_only=True)

    # ---- 課表 -----------------------------------------------------------------
    ws = wb["課表"]
    col = header_index(ws)
    need = ["週", "日期", "階段", "類型", "計畫km", "配速/心拍", "內容", "備註"]
    missing = [c for c in need if find_col(col, c) is None]
    if missing:
        sys.exit(f"課表 is missing columns {missing}; found {list(col)}")

    def cell(r, name):
        i = find_col(col, name)
        return r[i] if i is not None and i < len(r) else None

    sessions, results, unknown = [], [], set()
    phase_by_week: dict[int, str] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not isinstance(cell(r, "日期"), dt.datetime):
            continue
        date = cell(r, "日期").date().isoformat()
        label = text(cell(r, "類型"))
        kind = TYPE_MAP.get(label)
        if kind is None:
            unknown.add(label)
            continue
        wk = cell(r, "週")
        week_no = int(wk) if isinstance(wk, (int, float)) else None
        phase = text(cell(r, "階段")) or None
        if week_no is not None and phase and week_no not in phase_by_week:
            phase_by_week[week_no] = phase
        target = text(cell(r, "配速/心拍"))
        sessions.append({
            "date": date,
            "week_no": week_no,
            "phase": phase,
            "label": label,
            "type": kind,
            "distance_m": int(round(float(cell(r, "計畫km") or 0) * 1000)),
            "content": text(cell(r, "內容")) or None,
            "note": text(cell(r, "備註")) or None,
            "target": target or None,
            **parse_target(target),
        })
        # actuals, hand-typed; only rows with a verdict or a distance. 判定/提醒 columns are optional.
        if num(cell(r, "實際km")) is not None or text(cell(r, "判定")):
            reminders = [x.strip() for x in text(cell(r, "提醒")).split("｜") if x.strip()]
            results.append({
                "date": date,
                "distance_m": int(round((num(cell(r, "實際km")) or 0) * 1000)),
                "memo": text(cell(r, "心得")) or None,
                "moving_time_sec": time_sec(cell(r, "時間")),
                "pace_sec_km": pace_sec(cell(r, "平均配速")),
                "avg_hr": int(num(cell(r, "平均心拍"))) if num(cell(r, "平均心拍")) is not None else None,
                "max_hr": int(num(cell(r, "最大心拍"))) if num(cell(r, "最大心拍")) is not None else None,
                "avg_watts": int(num(cell(r, "平均功率W"))) if num(cell(r, "平均功率W")) is not None else None,
                "kcal": int(num(cell(r, "卡路里"))) if num(cell(r, "卡路里")) is not None else None,
                "weather": text(cell(r, "天氣")) or None,
                "shoes": text(cell(r, "鞋")) or None,
                "device": text(cell(r, "裝置")) or None,
                "verdict": text(cell(r, "判定")) or None,
                "reminders": reminders,
            })

    # ---- 週總覽 --------------------------------------------------------------
    # Two layouts seen: (週, 起始日, 結束日, 階段, …, 重點) and (週, 起始日, 計畫總km, …, 備註).
    # Missing end = start + 6 days; missing phase = the week's first session phase; focus = 重點 or 備註.
    wsw = wb["週總覽"]
    wcol = header_index(wsw)
    weeks = []
    for r in wsw.iter_rows(min_row=2, values_only=True):
        def wcell(*names):
            for name in names:
                i = find_col(wcol, name)
                if i is not None and i < len(r):
                    return r[i]
            return None

        wk, start = wcell("週"), wcell("起始日", "起始", "起", "開始")
        if not isinstance(wk, (int, float)) or not isinstance(start, dt.datetime):
            continue
        end_v = wcell("結束日", "結束", "迄")
        end = end_v.date() if isinstance(end_v, dt.datetime) else start.date() + dt.timedelta(days=6)
        focus = None
        for name in ("重點", "備註"):
            if find_col(wcol, name) is not None:
                focus = text(wcell(name)) or None
                break
        phase = text(wcell("階段"))
        weeks.append({
            "week_no": int(wk),
            "start": start.date().isoformat(),
            "end": end.isoformat(),
            "phase": phase or phase_by_week.get(int(wk)),
            "focus": focus,
        })
    # A derived end must not run into the next week.
    weeks.sort(key=lambda w: w["week_no"])
    for a, b in zip(weeks, weeks[1:]):
        if a["end"] >= b["start"]:
            a["end"] = (dt.date.fromisoformat(b["start"]) - dt.timedelta(days=1)).isoformat()

    if unknown:
        sys.exit(f"unmapped 類型 values: {sorted(unknown)} — add them to TYPE_MAP")

    sessions.sort(key=lambda s: s["date"])
    entry = {
        "user": user,
        "race_name": RACE_NAME,
        "race_date": RACE_DATE,
        "goal_time_sec": goal if goal is not None else DEFAULTS[user]["goal_time_sec"],
        "days_per_week": days if days is not None else DEFAULTS[user]["days_per_week"],
        "weeks": weeks,
        "sessions": sessions,
        "results": results,
    }

    # Merge: keep every other user's plan as it is, replace this user's.
    plans = json.loads(OUT.read_text(encoding="utf-8")) if OUT.exists() else []
    plans = [p for p in plans if p.get("user") != user] + [entry]
    plans.sort(key=lambda p: list(DEFAULTS).index(p["user"]) if p["user"] in DEFAULTS else 99)
    OUT.write_text(json.dumps(plans, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {OUT.relative_to(ROOT)} [{user}]: {len(weeks)} weeks, {len(sessions)} sessions "
          f"({sessions[0]['date']} → {sessions[-1]['date']}), {len(results)} results; "
          f"{len(plans)} plan(s) in file")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", nargs="?", default=DEFAULT_XLSX, type=pathlib.Path)
    ap.add_argument("--user", default="ww", choices=sorted(DEFAULTS), help="key in seed/seed.ts USERS")
    ap.add_argument("--goal", help="race goal H:MM, e.g. 5:30")
    ap.add_argument("--days", type=int, help="planned run days per week")
    a = ap.parse_args()
    main(a.xlsx, a.user, goal_sec(a.goal), a.days)
